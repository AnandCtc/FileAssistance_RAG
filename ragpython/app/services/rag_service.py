from langchain_chroma import Chroma

from langchain_openai import (
    OpenAIEmbeddings,
    ChatOpenAI
)

from langchain_text_splitters import RecursiveCharacterTextSplitter

from langchain_core.documents import Document

from langchain_core.messages import HumanMessage

from dotenv import load_dotenv

import os

# LOAD ENV VARIABLES

load_dotenv()

# VECTOR DB DIRECTORY

PERSIST_DIRECTORY = "db/chroma_db"

# OPENAI EMBEDDINGS

embedding_model = OpenAIEmbeddings(
    model="text-embedding-3-small"
)

# CHAT MODEL

chat_model = ChatOpenAI(
    model="gpt-3.5-turbo",
    temperature=0
)

# CHROMA DB

db = Chroma(
    persist_directory=PERSIST_DIRECTORY,
    embedding_function=embedding_model
)

# TEXT SPLITTER

text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=200
)

# PROCESS DOCUMENT

def process_document(file_path):

    print("STEP 6: Extracting text")

    text = extract_text(file_path)

    print("STEP 7: Chunking document")

    chunks = text_splitter.split_text(text)

    documents = [
        Document(page_content=chunk)
        for chunk in chunks
    ]

    print("STEP 8: Creating embeddings")

    db.add_documents(documents)

    print("STEP 9: Storing in ChromaDB")

    return True

# ASK QUESTION

def ask_question(question):

    print("STEP 14: Searching vector DB")

    retriever = db.as_retriever(
        search_kwargs={"k": 3}
    )

    relevant_docs = retriever.invoke(question)

    context = "\n\n".join([
        doc.page_content
        for doc in relevant_docs
    ])

    print("STEP 15: Sending context to GPT")

    prompt = f"""
    Answer the question based ONLY on the context below.

    Context:
    {context}

    Question:
    {question}
    """

    response = chat_model.invoke([
        HumanMessage(content=prompt)
    ])

    print("STEP 16: Returning answer")

    return response.content

# EXTRACT TEXT

def extract_text(file_path):

    extension = os.path.splitext(
        file_path
    )[1].lower()

    # PDF

    if extension == ".pdf":

        from pypdf import PdfReader

        reader = PdfReader(file_path)

        text = ""

        for page in reader.pages:
            text += page.extract_text() or ""

        return text

    # TXT

    elif extension == ".txt":

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as file:

            return file.read()

    # CSV

    elif extension == ".csv":

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as file:

            return file.read()

    # DOCX

    elif extension == ".docx":

        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)

        return "\n".join([
            para.text
            for para in doc.paragraphs
        ])

    else:

        raise Exception(
            f"Unsupported file type: {extension}"
        )